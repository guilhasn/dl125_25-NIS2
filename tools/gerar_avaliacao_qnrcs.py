"""
Gera o Excel de auto-avaliação de maturidade QNRCS v2.

Uso:
    python tools/gerar_avaliacao_qnrcs.py

Saída: templates/avaliacao-maturidade-qnrcs.xlsx (não encriptado — a
encriptação faz-se a seguir com aplicar_password.py).
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import RadarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

# --- estilos partilhados ------------------------------------------------

TITULO = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
SUBTITULO = Font(name="Calibri", size=11, italic=True, color="555555")
CAB = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TXT = Font(name="Calibri", size=11)
TXT_NEGRITO = Font(name="Calibri", size=11, bold=True)

AZUL = PatternFill("solid", fgColor="1F4E78")
AZUL_CLARO = PatternFill("solid", fgColor="D9E1F2")
CINZA = PatternFill("solid", fgColor="F2F2F2")
VERDE_CLARO = PatternFill("solid", fgColor="D9EAD3")
AMARELO_CLARO = PatternFill("solid", fgColor="FFF2CC")
VERMELHO_CLARO = PatternFill("solid", fgColor="F4CCCC")

LINHA = Side(border_style="thin", color="A6A6A6")
CAIXA = Border(left=LINHA, right=LINHA, top=LINHA, bottom=LINHA)

CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
ESQUERDA = Alignment(horizontal="left", vertical="center", wrap_text=True)
TOPO = Alignment(horizontal="left", vertical="top", wrap_text=True)

REPO = Path(__file__).resolve().parents[1]
DEST = REPO / "templates"

# --- dados dos 26 controlos ---------------------------------------------
# Cada tuplo: (código, descrição curta, mapeamento Anexo IV, aplicável Grupo B)

CONTROLOS_GR = [
    ("GR.CO-3", "Requisitos legais, regulamentares e contratuais geridos", "Transversal", "Sim (mínimo Básico)"),
    ("GR.GR-1", "Apetite e tolerância ao risco definidos", "Transversal · B.2", "Sim (mínimo Básico)"),
    ("GR.FR-1", "Órgãos de gestão compreendem funções e responsabilidades", "O.CRI · C.1", "Sim (mínimo Básico)"),
    ("GR.PP-1", "Políticas de cibersegurança estabelecidas", "Transversal", "Sim (mínimo Básico)"),
    ("GR.SP-1", "Supervisão da estratégia de cibersegurança", "E.1 supervisão", "Sim (mínimo Básico)"),
    ("GR.CA-1", "Políticas de gestão de risco da cadeia de abastecimento", "O.PSF · D.1", "Sim (mínimo Básico)"),
]

CONTROLOS_ID = [
    ("ID.GA-1", "Equipamentos e recursos físicos inventariados", "O.IAC · B.3", "Sim (mínimo Básico)"),
    ("ID.GA-5", "Activos classificados de acordo com a sua criticidade", "O.IAC · B.3", "Sim (mínimo Básico)"),
    ("ID.AR-1", "Vulnerabilidades dos activos identificadas e documentadas", "B.2 análise de risco", "Sim (mínimo Básico)"),
    ("ID.MC-1", "Melhorias identificadas através de avaliações", "E.2 documentação", "Sim (mínimo Básico)"),
]

CONTROLOS_PR = [
    ("PR.GA-1", "Ciclos de vida de gestão de identidades definidos", "O.GAP · D.3", "Sim (mínimo Básico)"),
    ("PR.GA-3", "Mecanismos de autenticação definidos (incluindo MFA)", "T.AM · D.3", "Sim (mínimo Básico)"),
    ("PR.FC-1", "Pessoal sensibilizado e formado em cibersegurança", "H.PF · D.4", "Sim (mínimo Básico)"),
    ("PR.SD-1", "Confidencialidade, integridade e disponibilidade dos dados protegida", "T.PEW · D.2", "Sim (mínimo Básico)"),
    ("PR.SD-5", "Cópias de segurança realizadas, mantidas e testadas", "T.CS · D.2", "Sim (mínimo Básico)"),
    ("PR.SP-1", "Configuração base de redes e sistemas criada e mantida", "T.PAS · T.AS · D.3", "Sim (mínimo Básico)"),
    ("PR.RI-1", "Integridade das redes de comunicações protegida", "T.AS · D.2", "Sim (mínimo Básico)"),
]

CONTROLOS_DE = [
    ("DE.MC-1", "Redes e sistemas monitorizados para detectar incidentes", "T.MA", "Sim (mínimo Básico)"),
    ("DE.AE-1", "Eventos detectados são analisados", "O.GEC", "Sim (mínimo Básico)"),
    ("DE.AE-3", "Impacto estimado e âmbito dos eventos adversos compreendidos", "C.2 prazos", "Sim (mínimo Básico)"),
]

CONTROLOS_RS = [
    ("RS.GI-1", "Plano de Resposta a Incidentes executado em coordenação", "O.CRI · C.1", "Sim (mínimo Básico)"),
    ("RS.AI-1", "Análise forense para determinar o que ocorreu no incidente", "C.6", "Sim (mínimo Básico)"),
    ("RS.NC-1", "Partes interessadas notificadas dos incidentes", "art. 42.º · C.2/C.3", "Sim (mínimo Básico)"),
    ("RS.MI-1", "Incidentes são contidos", "C.6", "Sim (mínimo Básico)"),
]

CONTROLOS_RC = [
    ("RC.PR-1", "Plano de recuperação seguido durante/após incidente", "T.CS · D.2", "Sim (mínimo Básico)"),
    ("RC.CO-1", "Actividades de recuperação comunicadas a partes interessadas", "art. 48.º · C.4", "Sim (mínimo Básico)"),
]

OBJECTIVOS = [
    ("GR — Gerir", CONTROLOS_GR),
    ("ID — Identificar", CONTROLOS_ID),
    ("PR — Proteger", CONTROLOS_PR),
    ("DE — Detectar", CONTROLOS_DE),
    ("RS — Responder", CONTROLOS_RS),
    ("RC — Recuperar", CONTROLOS_RC),
]

# --- abas ----------------------------------------------------------------

def aba_instrucoes(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "Instruções"

    ws.column_dimensions["A"].width = 110
    ws.row_dimensions[1].height = 30

    ws["A1"] = "Auto-avaliação de maturidade QNRCS v2 — Instruções"
    ws["A1"].font = TITULO
    ws["A1"].fill = AZUL
    ws["A1"].alignment = CENTRO

    linhas = [
        ("Base legal", "Art. 23.º + Anexo I do Aviso n.º 5146/2026/2 (Quadro Nacional de Referência para a Cibersegurança v2)."),
        ("Objectivo", "Medir o grau de maturidade da câmara em 26 controlos de cibersegurança seleccionados do QNRCS v2, alinhados com as medidas do Anexo IV. O resultado alimenta o roadmap E.3."),
        ("Quem preenche", "Equipa de 2-3 pessoas: TIC + DPO + RH/jurídico, idealmente em 2 sessões de trabalho (~1h cada)."),
        ("Os 6 objectivos do QNRCS v2", "Gerir (GR) · Identificar (ID) · Proteger (PR) · Detectar (DE) · Responder (RS) · Recuperar (RC). Cada um tem uma aba dedicada."),
        ("Os 3 níveis cumulativos", "Básico ⊂ Substancial ⊂ Elevado. Para câmara Grupo B o mínimo é Básico em todos os controlos. Substancial/Elevado são aspiracionais (ou aplicáveis a SMAS/empresa municipal essencial)."),
        ("Como preencher cada controlo", "Coluna C 'Estado actual': Não / Parcial / Sim · Coluna D 'Nível atingido': Não cumpre / Básico / Substancial / Elevado · Coluna G 'Evidência existente': documento, template ou registo que comprova (obrigatório se C=Sim)."),
        ("Dashboard", "Após preencher as 6 abas, a aba Dashboard mostra automaticamente o perfil de maturidade por objectivo (% e gráfico radar) e a lista dos controlos em défice a priorizar no roadmap E.3."),
        ("Periodicidade", "Repetir anualmente, idealmente em paralelo com a revisão do dossier mínimo de conformidade (E.2)."),
    ]

    linha_actual = 3
    for titulo, texto in linhas:
        ws.cell(row=linha_actual, column=1, value=titulo).font = TXT_NEGRITO
        ws.cell(row=linha_actual, column=1).fill = AZUL_CLARO
        ws.cell(row=linha_actual, column=1).alignment = ESQUERDA
        ws.cell(row=linha_actual, column=1).border = CAIXA
        ws.row_dimensions[linha_actual].height = 22

        ws.cell(row=linha_actual + 1, column=1, value=texto).font = TXT
        ws.cell(row=linha_actual + 1, column=1).alignment = TOPO
        ws.cell(row=linha_actual + 1, column=1).border = CAIXA
        ws.row_dimensions[linha_actual + 1].height = 55

        linha_actual += 2

def aba_objectivo(wb: Workbook, titulo: str, controlos: list) -> None:
    """Cria uma aba com a lista de controlos de um objectivo."""
    sigla = titulo.split(" — ")[0]
    ws = wb.create_sheet(title=sigla)

    # Larguras
    ws.column_dimensions["A"].width = 12  # Código
    ws.column_dimensions["B"].width = 60  # Descrição
    ws.column_dimensions["C"].width = 14  # Estado
    ws.column_dimensions["D"].width = 18  # Nível atingido
    ws.column_dimensions["E"].width = 30  # Mapeamento
    ws.column_dimensions["F"].width = 22  # Aplicável Grupo B
    ws.column_dimensions["G"].width = 40  # Evidência
    ws.column_dimensions["H"].width = 30  # Notas

    # Faixa de título
    ws.merge_cells("A1:H1")
    ws["A1"] = f"Objectivo {titulo}"
    ws["A1"].font = TITULO
    ws["A1"].fill = AZUL
    ws["A1"].alignment = CENTRO
    ws.row_dimensions[1].height = 28

    # Cabeçalhos
    cabecalhos = [
        "Código",
        "Descrição do controlo",
        "Estado actual",
        "Nível atingido",
        "Mapeamento Anexo IV",
        "Aplicável Grupo B?",
        "Evidência existente",
        "Notas",
    ]
    for col, texto in enumerate(cabecalhos, start=1):
        c = ws.cell(row=2, column=col, value=texto)
        c.font = CAB
        c.fill = AZUL
        c.alignment = CENTRO
        c.border = CAIXA
    ws.row_dimensions[2].height = 35

    # Linhas de controlos
    primeira_linha_dados = 3
    for i, (codigo, descricao, mapeamento, aplicavel) in enumerate(controlos, start=primeira_linha_dados):
        ws.cell(row=i, column=1, value=codigo).font = TXT_NEGRITO
        ws.cell(row=i, column=2, value=descricao).font = TXT
        ws.cell(row=i, column=3, value="")  # Estado (utilizador preenche)
        ws.cell(row=i, column=4, value="")  # Nível atingido (utilizador preenche)
        ws.cell(row=i, column=5, value=mapeamento).font = TXT
        ws.cell(row=i, column=6, value=aplicavel).font = TXT
        ws.cell(row=i, column=7, value="")  # Evidência (utilizador preenche)
        ws.cell(row=i, column=8, value="")  # Notas

        for col in range(1, 9):
            ws.cell(row=i, column=col).border = CAIXA
            ws.cell(row=i, column=col).alignment = TOPO if col in (2, 5, 7, 8) else CENTRO

        ws.row_dimensions[i].height = 50

    ultima_linha_dados = primeira_linha_dados + len(controlos) - 1

    # Data validation — Estado (coluna C)
    dv_estado = DataValidation(
        type="list", formula1='"Não,Parcial,Sim"', allow_blank=True
    )
    dv_estado.add(f"C{primeira_linha_dados}:C{ultima_linha_dados}")
    ws.add_data_validation(dv_estado)

    # Data validation — Nível atingido (coluna D)
    dv_nivel = DataValidation(
        type="list",
        formula1='"Não cumpre,Básico,Substancial,Elevado"',
        allow_blank=True,
    )
    dv_nivel.add(f"D{primeira_linha_dados}:D{ultima_linha_dados}")
    ws.add_data_validation(dv_nivel)

    # Formatação condicional — Nível atingido
    intervalo_nivel = f"D{primeira_linha_dados}:D{ultima_linha_dados}"
    ws.conditional_formatting.add(
        intervalo_nivel,
        CellIsRule(operator="equal", formula=['"Não cumpre"'], fill=VERMELHO_CLARO),
    )
    ws.conditional_formatting.add(
        intervalo_nivel,
        CellIsRule(operator="equal", formula=['"Básico"'], fill=AMARELO_CLARO),
    )
    ws.conditional_formatting.add(
        intervalo_nivel,
        CellIsRule(operator="equal", formula=['"Substancial"'], fill=VERDE_CLARO),
    )
    ws.conditional_formatting.add(
        intervalo_nivel,
        CellIsRule(operator="equal", formula=['"Elevado"'], fill=VERDE_CLARO),
    )


# --- main ---------------------------------------------------------------

def main() -> None:
    DEST.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    aba_instrucoes(wb)
    for titulo, controlos in OBJECTIVOS:
        aba_objectivo(wb, titulo, controlos)
    # task 4 acrescenta o dashboard aqui

    destino = DEST / "avaliacao-maturidade-qnrcs.xlsx"
    wb.save(destino)
    print(f"OK -> {destino}")


if __name__ == "__main__":
    main()
